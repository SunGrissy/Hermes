# -*- coding: utf-8 -*-
"""
本地 Ollama（默认 gemma4）+ game-review 技能要点，对策划方案文档做审核，输出 Markdown 报告。

支持输入：.md / .txt / .pdf / .docx

用法（PowerShell）:
  cd d:\\MyAgents
  $env:GAME_REVIEW_LLM_API_BASE = "http://127.0.0.1:11434/v1"
  $env:GAME_REVIEW_LLM_MODEL = "gemma4:latest"
  py scripts/game_review_ollama.py 方案.md -o 审核报告.md

  py scripts/game_review_ollama.py 方案.pdf -o out.md
  py scripts/game_review_ollama.py 方案.docx --scheme-type 活动方案

  # 多轮模式：每个引擎独立调用，再汇总（更深但更慢）
  py scripts/game_review_ollama.py 方案.pdf --multi-pass -o out.md

依赖:
  pip install openai pypdf python-docx

前置: Ollama 已启动并已拉取模型，例如: ollama pull gemma4:latest
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import sys
from pathlib import Path

# Windows 终端 UTF-8
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

def _default_env() -> None:
    os.environ.setdefault("GAME_REVIEW_LLM_API_BASE", "http://127.0.0.1:11434/v1")
    os.environ.setdefault("GAME_REVIEW_LLM_API_KEY", "ollama")
    os.environ.setdefault("GAME_REVIEW_LLM_MODEL", "gemma4:latest")


def _load_text_from_file(path: Path) -> str:
    suf = path.suffix.lower()
    if suf in (".md", ".markdown", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")

    if suf == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise SystemExit(
                "缺少依赖 pypdf，请执行: pip install pypdf"
            ) from e
        reader = PdfReader(str(path))
        parts: list[str] = []
        for page in reader.pages:
            t = page.extract_text() or ""
            parts.append(t)
        return "\n".join(parts)

    if suf == ".docx":
        try:
            from docx import Document
        except ImportError as e:
            raise SystemExit(
                "缺少依赖 python-docx，请执行: pip install python-docx"
            ) from e
        doc = Document(str(path))
        lines: list[str] = []
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines)

    raise SystemExit(f"不支持的扩展名: {suf}（支持 .md .txt .pdf .docx）")


def _truncate(text: str, max_chars: int) -> tuple[str, bool]:
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars] + "\n\n[… 正文已截断，仅保留前 max_chars 字符 …]\n", True


def _build_system_prompt(scheme_type: str) -> str:
    """嵌入 game-review 核心判据（与 SKILL references 对齐的浓缩版）。"""
    return f"""你是游戏策划方案审核助手，遵循「制作人方法论」与 game-review 六引擎框架。
当前方案类型（用户指定）：{scheme_type}

## 审核纪律（必须遵守）

### 宁严勿松
- 对每个引擎，**先找问题再确认优点**。不要默认给 ✅ 再找例外。
- 如果一个模块的核心参数/数值/规则未填写，该模块的服务性**不得**标为 ✅——没有参数等于没有设计。
- 不要因为"方向好"就拉高分数；方向好但细节空 = 框架存在但不可执行。

### 矛盾检测
逐一检查文档中不同模块之间的逻辑是否自洽：
- 规则 A 说的和规则 B 说的是否矛盾？
- 某个数值/机制在不同场景下的表现是否冲突？
- 设计意图和实际机制是否对得上？
发现矛盾时标为 P1 并具体说明两条规则各在文档哪里。

## 价值排序（默认）
公平感 > 社交生态 > 爽感 > 留存 > 付费

## 三条红线（摘要）
- 核心循环不为 KPI 让步
- 不强迫玩家
- 不做冗余系统

## 五大误区
1. 堆系统当内容（系统多≠内容丰富）
2. 用数值手段解决体验问题（加数字≠解决感受）
3. 过度参考竞品（像≠好）
4. 前期过度披露（一眼看完≠信息透明）
5. 设计给自己玩（作者脑补了完整体验但读者/玩家无法还原）

## 六引擎检查要点（须逐引擎输出小节）
1. **Purpose（设计目的）**：能否一句话说清目的；模块是否都服务该目的；目标人群是否明确。
2. **Rhythm（情绪节奏）**：画出情绪曲线（阶段→情绪→类型），检查是否存在**情绪平原**（长时间无波动）；宏观节奏是否合理。
3. **RedLine（红线与误区）**：是否触碰三条红线或命中五大误区。
4. **Value（价值冲突）**：行为是否违反上述价值排序；**不同玩家群体之间是否存在公平感冲突**。
5. **Segment（玩家分层）**：大R/中坚/免费/低活跃分别受到什么影响；中坚层是否被保护；是否有分层被完全忽略。
6. **Economy（经济健康）**：通胀/重置/追赶/回收等是否有明显风险（若文档几乎无数值可写「本稿侧重非数值，经济项标为待验证」）。

## 引擎裁剪（供你自检是否弱化某些引擎）
- 完整策划案：六引擎全开
- 单系统设计：Purpose, Rhythm, RedLine 必做；Value/Segment/Economy 视内容选做
- 数值方案：Purpose, RedLine, Economy 必做
- 活动/运营方案：Purpose, RedLine, Value 必做；其余视内容
- UI/交互：Purpose, Rhythm, RedLine

## 输出要求（严格使用 Markdown，勿用代码块包裹整篇）
按下列结构输出（可自拟小节标题，但必须包含这些块）：

# [方案名称或文件名] 审核报告

**审核日期**：{dt.date.today().isoformat()}
**方案类型**：{scheme_type}
**启用引擎**：列出实际着墨的引擎

## 一、30秒速判
- 设计目的（一句话）
- 宏观适配：✅/⚠️/❌
- 微观目标感：简述

## 二、引擎检查结果
为每个启用的引擎写「### X 引擎」小节，内含：检查项表格或列表、问题清单（标 P0/P1/P2）、维度评分（10分制）、必要时一条 Gate/Flag/Sense 式「决策问题」供制作人判断。

## 三、制作人决策记录（模板）
用表格列出待制作人回答的决策问题（类型 Gate/Flag/Sense | 问题 | 可选选项），制作人未在场则留空列。每个决策问题**必须提供 A/B/C 三个选项**：A=明确处理方案，B=替代方案或不处理的理由，C=不确定/需要更多数据再判断。不接受是/否二选一。

## 四、问题汇总
- P0 / P1 / P2 分级列表
- 待验证项

## 五、综合评分
表格：各引擎维度分 + 综合分（10分制）

## 六、文档完整性备注（附录）
六引擎审完后，最后扫一遍：文档中是否存在大面积空白的数值参数、未填写的决策项、或无锚定值的核心规则？如果有，在此列出，标注对下游（UX/数值/技术）的阻塞影响。如果文档填写完整，写"无明显缺项"即可。此章节不影响上方各引擎评分——引擎只审已写的内容。

语言：简体中文。基于文档内容作判断，不要编造文档中不存在的事实；信息不足处明确写「文档未提及/待补充」。
"""


_ENGINE_DEFS: list[dict[str, str]] = [
    {
        "code": "Purpose",
        "name": "设计目的检查",
        "prompt": """你是 Purpose 引擎。只检查以下四项，逐项给出表格（检查项|结论✅⚠️❌|说明）：
1. 设计目的清晰度：能否一句话说清核心目的？需要两段话才能解释 = 没想明白。
2. 模块服务性：方案中每个模块/子系统是否都服务核心目的？"为了有而有"的模块标 ⚠️。如果模块的核心参数未填写，不得标 ✅。
3. 目标人群明确性：是否明确了服务谁？
4. 宏观适配：在整体游戏框架中位置是否合理？
检查完后输出问题清单（P0/P1/P2）、维度评分（10分制）。如有 P0 问题，生成一条 Gate 决策问题（A/B/C 三选项）。""",
    },
    {
        "code": "Rhythm",
        "name": "情绪节奏分析",
        "prompt": """你是 Rhythm 引擎。执行以下检查：
1. 画出情绪曲线表格（阶段 | 情绪 | 类型：压力期/波动期/高峰/释放等）。
2. 检查是否存在**情绪平原**（某阶段超过3天却无情绪波动设计）。如有，标为 ⚠️ 并说明。
3. 宏观节奏合理性：周期长度、递进结构是否合理。
4. 关键衔接点：阶段之间的过渡是否平滑，是否有情绪断崖。
输出检查项表格、问题清单（P0/P1/P2）、维度评分（10分制）。如发现情绪平原，生成一条 Gate 决策问题（A/B/C 三选项）。""",
    },
    {
        "code": "RedLine",
        "name": "红线与误区扫描",
        "prompt": """你是 RedLine 引擎。检查两大类：
**红线检查**（逐项表格）：
- 核心循环是否被 KPI 污染
- 是否存在强迫设计（玩家无法选择不参与）
- 是否存在冗余系统（有路径但无实际体验价值）
**五大误区检查**（逐项表格）：
1. 堆系统当内容（系统多≠内容丰富）
2. 用数值手段解决体验问题（加数字≠解决感受）——注意边界：用数值制造心理压力是否属于此误区？
3. 过度参考竞品
4. 前期过度披露
5. 设计给自己玩（作者脑补了体验但读者/玩家无法还原——文档中大量空白参数是此误区的信号）
输出问题清单（P0/P1/P2）、维度评分（10分制）。""",
    },
    {
        "code": "Value",
        "name": "价值冲突检测",
        "prompt": """你是 Value 引擎。默认价值排序：公平感 > 社交生态 > 爽感 > 留存 > 付费。
检查以下内容（逐项表格）：
1. 方案中是否存在机制违反上述排序？（如付费优势压过公平感）
2. 不同玩家群体之间是否存在公平感冲突？（如不同分组/赛区的难度不一致）
3. 社交生态影响：排名/荣誉/展示机制对社群的正负面效应。
4. 付费设计合理性：付费点与核心体验的挂钩方式。
特别注意：如果文档中存在乘法加成（如×3）且无追赶手段，这是公平感冲突的典型信号。
输出问题清单（P0/P1/P2）、维度评分（10分制）、必要时生成 Flag 决策问题（A/B/C 三选项）。""",
    },
    {
        "code": "Segment",
        "name": "玩家分层评估",
        "prompt": """你是 Segment 引擎。按以下分层逐一评估影响（表格：分层|影响✅⚠️❌|说明）：
- 大R（鲸鱼/高付费）：体验是否匹配投入
- 中坚层（海豚/高活跃/强社交）：是否被保护，是否能承上启下——这是最重要的分层
- 免费/低付费层：是否有事干，是否感到被孤立或挫败
- 高攻击性玩家：是否有限制机制
- 低活跃玩家：是否有追赶/回归路径
如果某个分层在文档中**完全未被提及**，标为 ❌ 并注明"文档忽略此分层"。
检查是否存在：门槛把某个分层完全排除在外但又给他们曝光 → 看得见进不去 = 挫败感。
输出问题清单（P0/P1/P2）、维度评分（10分制）。""",
    },
    {
        "code": "Economy",
        "name": "经济系统健康度",
        "prompt": """你是 Economy 引擎。检查以下内容（逐项表格）：
1. 通胀风险：奖励/产出是否有周期性重置？重复发放是否会造成资源池通胀？
2. 重置策略：付费养成是否被重置？非付费养成的重置是否合理？
3. 追赶机制：后来者/低资源玩家是否有追赶路径？
4. 资源回收闭环：产出的资源是否有消耗出口？是否会"沉底"？
5. 付费/赛季分离：付费维度和赛季/活动维度是否清晰分开？
如果文档几乎无数值内容，写「本稿侧重非数值设计，经济项标为待验证」，仍给出基于定性描述的判断。
输出问题清单（P0/P1/P2）、维度评分（10分制）。""",
    },
]


def _build_engine_system(engine: dict[str, str], scheme_type: str) -> str:
    return f"""你是游戏策划方案审核助手的「{engine['name']}」模块（{engine['code']} 引擎）。
当前方案类型：{scheme_type}

## 审核纪律
- 宁严勿松：先找问题再确认优点。参数未填写的模块不得标 ✅。
- 基于文档内容判断，不编造文档中不存在的事实；信息不足写「文档未提及」。
- 如果发现文档内不同模块之间存在逻辑矛盾，标为 P1 并说明。
- 决策问题必须提供 A/B/C 三选项。
- 语言：简体中文。

## 你的检查任务
{engine['prompt']}
"""


def _build_aggregate_system(scheme_type: str) -> str:
    return f"""你是游戏策划方案审核的汇总模块。下面是 6 个引擎的独立检查结果。
请将它们整合为一份完整审核报告，格式如下：

# [方案名称] 审核报告

**审核日期**：{dt.date.today().isoformat()}
**方案类型**：{scheme_type}
**审核模式**：多轮独立引擎
**启用引擎**：列出实际着墨的引擎

## 一、30秒速判
- 设计目的（一句话，从 Purpose 引擎结果提炼）
- 宏观适配：✅/⚠️/❌
- 微观目标感：简述

## 二、引擎检查结果
直接搬运各引擎输出，每个引擎一个 ### 小节。保留原有表格、评分、问题清单和决策问题。如发现多个引擎指向同一问题，在该引擎处标注「与 X 引擎交叉」。

## 三、制作人决策记录（模板）
汇总所有引擎产生的决策问题到一张表（类型 | 引擎来源 | 问题 | 选项A | 选项B | 选项C）。

## 四、问题汇总
合并去重所有引擎的问题，按 P0/P1/P2 分级。

## 五、综合评分
表格：各引擎维度分 + 综合分（10分制）。综合分 = 各引擎简单平均，不要人为拉高。

## 六、文档完整性备注（附录）
最后扫一遍：文档中是否存在大面积空白的数值参数、未填写的决策项？如有列出，标注对下游的阻塞影响。此章节不影响上方引擎评分。

语言：简体中文。"""


def _run_multi_pass(
    body: str,
    in_name: str,
    scheme_type: str,
    *,
    base: str,
    api_key: str,
    model: str,
    timeout: int,
) -> str:
    user_doc = f"**源文件**：{in_name}\n\n---\n\n{body}"
    engine_results: list[str] = []
    total = len(_ENGINE_DEFS)

    for i, eng in enumerate(_ENGINE_DEFS, 1):
        print(f"  [{i}/{total}] {eng['code']}({eng['name']})...", file=sys.stderr)
        sys_prompt = _build_engine_system(eng, scheme_type)
        result = _call_openai_compatible(
            sys_prompt,
            f"请对以下文档执行 {eng['code']} 引擎检查。\n\n{user_doc}",
            base=base,
            api_key=api_key,
            model=model,
            timeout=timeout,
        )
        engine_results.append(f"### {eng['code']}（{eng['name']}）\n\n{result}")

    combined = "\n\n---\n\n".join(engine_results)
    print(f"  [汇总] 正在生成最终报告...", file=sys.stderr)
    agg_system = _build_aggregate_system(scheme_type)
    report = _call_openai_compatible(
        agg_system,
        f"以下是 6 个引擎的独立检查结果，请汇总为完整报告。\n\n{combined}",
        base=base,
        api_key=api_key,
        model=model,
        timeout=timeout,
    )
    return report


def _call_openai_compatible(
    system: str,
    user: str,
    *,
    base: str,
    api_key: str,
    model: str,
    timeout: int,
) -> str:
    try:
        from openai import OpenAI
    except ImportError as e:
        raise SystemExit("缺少 openai，请执行: pip install openai") from e

    client = OpenAI(base_url=base.rstrip("/"), api_key=api_key, timeout=timeout)
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": user},
        ],
        temperature=0.3,
    )
    choice = resp.choices[0].message
    content = (choice.content or "").strip()
    if not content:
        raise RuntimeError("模型返回空内容")
    return content


def main() -> int:
    _default_env()
    p = argparse.ArgumentParser(
        description="Ollama(gemma4 等) + game-review 要点，生成策划方案审核 Markdown"
    )
    p.add_argument(
        "input_file",
        type=Path,
        help="输入文件：.md .txt .pdf .docx",
    )
    p.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="输出 .md 路径；默认 输入文件名_game_review_日期时间.md",
    )
    p.add_argument(
        "--scheme-type",
        default="完整策划案",
        help="方案类型，用于引擎裁剪提示（默认：完整策划案）",
    )
    p.add_argument(
        "--max-chars",
        type=int,
        default=100_000,
        help="正文最大字符数，超出截断（默认 100000）",
    )
    p.add_argument(
        "--timeout",
        type=int,
        default=600,
        help="API 超时秒数（本地大模型可能较慢，默认 600）",
    )
    p.add_argument(
        "--multi-pass",
        action="store_true",
        help="多轮模式：每个引擎独立调用再汇总（更深但更慢，约 6+1 轮）",
    )
    p.add_argument(
        "--dry-run",
        action="store_true",
        help="只提取正文并打印长度，不调模型",
    )
    args = p.parse_args()

    in_path: Path = args.input_file.resolve()
    if not in_path.is_file():
        print(f"文件不存在: {in_path}", file=sys.stderr)
        return 1

    raw = _load_text_from_file(in_path)
    if len(raw.strip()) < 20:
        print("提取的正文过短，请检查 PDF/文档是否扫描版或加密。", file=sys.stderr)
        return 1

    body, truncated = _truncate(raw, args.max_chars)

    base = os.environ.get("GAME_REVIEW_LLM_API_BASE", "").rstrip("/")
    key = os.environ.get("GAME_REVIEW_LLM_API_KEY", "ollama")
    model = os.environ.get("GAME_REVIEW_LLM_MODEL", "gemma4:latest")

    if args.dry_run:
        print(f"正文长度: {len(raw)} 字符（截断后 {len(body)}） truncated={truncated}")
        print(f"API_BASE={base} MODEL={model}")
        return 0

    mode = "multi-pass" if args.multi_pass else "single"
    print(f"正在调用本地模型（{mode}），请稍候…", file=sys.stderr)
    print(f"  GAME_REVIEW_LLM_API_BASE={base}", file=sys.stderr)
    print(f"  GAME_REVIEW_LLM_MODEL={model}", file=sys.stderr)

    try:
        if args.multi_pass:
            report = _run_multi_pass(
                body,
                in_path.name,
                args.scheme_type,
                base=base,
                api_key=key,
                model=model,
                timeout=args.timeout,
            )
        else:
            system = _build_system_prompt(args.scheme_type)
            user = f"""请审核以下策划方案文档。

**源文件**：{in_path.name}
**说明**：以下为全文或截断片段。

---

{body}
"""
            report = _call_openai_compatible(
                system,
                user,
                base=base,
                api_key=key,
                model=model,
                timeout=args.timeout,
            )
    except Exception as e:
        print(f"调用失败: {e}", file=sys.stderr)
        return 1

    out = args.output
    if out is None:
        stem = in_path.stem
        ts = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        out = in_path.parent / f"{stem}_game_review_{ts}.md"
    else:
        out = out.resolve()

    header = f"""<!-- game-review-ollama | mode={mode} | model={model} | source={in_path.name} | truncated={truncated} -->
"""
    out.write_text(header + report + "\n", encoding="utf-8")
    print(str(out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
